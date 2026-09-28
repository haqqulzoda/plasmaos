"""W8 deterministic sealing, projection, staleness, and private artifact authority."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
from typing import Any
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfgen.canvas import Canvas
from xml.sax.saxutils import escape
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.private_storage import resolve_private_storage_key
from app.models.all_models import Organization, OrganizationPursuit, Proposal, Tender
from app.models.base import MembershipState, PursuitOrigin
from app.models.candidate_retrieval import CVVersion, ProjectReference
from app.models.private_documents import PursuitTenderContext
from app.models.proposal_evidence import (
    ProposalEvidenceArtifact,
    ProposalEvidencePack,
    ProposalEvidencePackItem,
    PursuitProposalWorkspace,
)
from app.models.pursuit_analysis import AnalysisCompanySnapshot, AnalysisPackItem
from app.models.team_scenarios import TeamScenario
from app.models.tenancy import Membership
from app.schemas.proposal_evidence import (
    ProposalEvidenceArtifactResponse,
    ProposalEvidenceExportRequest,
    ProposalEvidenceItemResponse,
    ProposalEvidencePackResponse,
    ProposalEvidenceSealRequest,
    PursuitProposalWorkspaceResponse,
)
from app.services.pursuit_analysis import get_analysis_run
from app.services.team_scenarios import (
    TeamScenarioEligibilityError,
    _project_scenarios,
    get_proposal_handoff,
)


PACK_SCHEMA_VERSION = "w8.proposal-evidence-pack.v1"
EXPORT_GENERATOR_VERSION = "w8.deterministic-export.v1"


class ProposalEvidenceError(ValueError):
    pass


class ProposalEvidenceNotFoundError(ProposalEvidenceError):
    pass


class ProposalEvidenceEligibilityError(ProposalEvidenceError):
    pass


def _json_default(value: Any) -> str:
    if isinstance(value, (UUID, datetime, date, Decimal)):
        return str(value)
    if hasattr(value, "value"):
        return str(value.value)
    raise TypeError(f"Unsupported manifest value: {type(value).__name__}")


def _jsonable(value: Any) -> Any:
    return json.loads(json.dumps(value, default=_json_default, ensure_ascii=False))


def _canonical(value: Any) -> bytes:
    return json.dumps(
        _jsonable(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")


async def _active_membership(db: AsyncSession, organization_id: UUID, membership_id: UUID) -> Membership:
    membership = await db.scalar(select(Membership).where(
        Membership.id == membership_id,
        Membership.organization_id == organization_id,
        Membership.state == MembershipState.ACTIVE,
    ))
    if membership is None:
        raise ProposalEvidenceEligibilityError("An active Organization Membership is required")
    return membership


async def _lock_sealing_authorities(db: AsyncSession) -> None:
    """Prevent authoritative append/change races until the pack transaction commits."""
    if db.bind is None or db.bind.dialect.name != "postgresql":
        return
    await db.execute(text(
        "LOCK TABLE memberships, pursuit_analysis_runs, pursuit_analysis_review_assertions, "
        "candidate_review_decisions, candidate_availability_facts, candidate_interest_facts, "
        "candidate_participation_decisions, team_scenario_revisions, team_scenario_decisions "
        "IN SHARE MODE"
    ))
    await db.execute(text(
        "LOCK TABLE pursuit_proposal_workspaces, proposal_evidence_packs IN SHARE ROW EXCLUSIVE MODE"
    ))


def _manifest_item(
    *, category: str, authority: str, identity: UUID | str, provenance: str,
    purpose: str, payload: dict[str, Any], review_state: str | None = None,
    evidence_state: str | None = None, requirement_id: UUID | None = None,
    position_id: UUID | None = None, gap_id: UUID | None = None,
    source_sha256: str | None = None, source_version: str | int | None = None,
) -> dict[str, Any]:
    return {
        "category": category,
        "source_authority_type": authority,
        "source_identity": str(identity),
        "provenance": provenance,
        "review_state": review_state,
        "evidence_state": evidence_state,
        "purpose": purpose,
        "requirement_id": requirement_id,
        "position_id": position_id,
        "gap_id": gap_id,
        "source_sha256": source_sha256,
        "source_version": str(source_version) if source_version is not None else None,
        "payload": _jsonable(payload),
    }


async def _pursuit_context_item(
    db: AsyncSession, pursuit: OrganizationPursuit,
) -> tuple[dict[str, Any], str]:
    if pursuit.origin == PursuitOrigin.SOURCE and pursuit.source_tender_id:
        tender = await db.get(Tender, pursuit.source_tender_id)
        if tender is None:
            raise ProposalEvidenceEligibilityError("Source Pursuit tender is unavailable")
        title = tender.title
        payload = {
            "pursuit_id": pursuit.id, "origin": "SOURCE", "title": tender.title,
            "buyer": tender.buyer, "country": tender.country, "reference": tender.external_id,
            "deadline": tender.deadline, "source_system": tender.source_system,
            "source_url": tender.source_url, "procurement_method": tender.procurement_method,
            "world_bank_boundary": {
                "project_context_is_context_only": True,
                "project_leadership_excluded_from_team_and_evidence": True,
            },
        }
        return _manifest_item(
            category="PURSUIT_CONTEXT", authority="Tender", identity=tender.id,
            provenance="SHARED_SOURCE", purpose="Pursuit identity and source context; excludes commercial price",
            payload=payload,
        ), title
    context = await db.get(PursuitTenderContext, pursuit.id)
    title = (context.title if context else None) or "Uploaded pursuit"
    payload = {
        "pursuit_id": pursuit.id, "origin": "UPLOAD", "title": title,
        "buyer": context.buyer if context else None,
        "declared_funder": context.declared_funder if context else None,
        "country": context.country if context else None,
        "reference": context.reference if context else None,
        "procurement_stage": context.procurement_stage if context else None,
        "deadline": context.external_deadline if context else None,
        "deadline_timezone": context.deadline_timezone if context else None,
        "source_url": context.source_url if context else None,
        "confirmed_fields": context.confirmed_fields if context else {},
    }
    return _manifest_item(
        category="PURSUIT_CONTEXT", authority="PursuitTenderContext", identity=pursuit.id,
        provenance="ORGANIZATION_PRIVATE_UPLOAD", purpose="Customer-reviewed uploaded Pursuit context; excludes commercial price",
        payload=payload, review_state="CONFIRMED" if context and context.confirmed_fields else "PROVISIONAL",
    ), title


def _selected_evidence_ids(handoff: Any) -> tuple[set[UUID], set[UUID]]:
    references: set[UUID] = set()
    cvs: set[UUID] = set()
    for participant in handoff.participants:
        for contribution in participant.contributions:
            for identity in contribution.evidence_identities:
                try:
                    value = UUID(str(identity.get("id")))
                except (ValueError, TypeError, AttributeError):
                    continue
                if identity.get("type") == "PROJECT_REFERENCE":
                    references.add(value)
                elif identity.get("type") == "CV_VERSION":
                    cvs.add(value)
    return references, cvs


async def _build_manifest(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID,
    scenario: TeamScenario, handoff: Any,
) -> tuple[list[dict[str, Any]], str, dict[str, int]]:
    pursuit = await db.scalar(select(OrganizationPursuit).where(
        OrganizationPursuit.id == pursuit_id,
        OrganizationPursuit.organization_id == organization_id,
    ))
    if pursuit is None:
        raise ProposalEvidenceNotFoundError("Pursuit not found")
    analysis = await get_analysis_run(
        db, organization_id=organization_id, pursuit_id=pursuit_id, run_id=handoff.analysis_run_id,
    )
    if analysis is None or analysis.analysis_pack_id != handoff.analysis_pack_id or analysis.inputs_changed:
        raise ProposalEvidenceEligibilityError("The exact W4 analysis is no longer current")

    items: list[dict[str, Any]] = []
    context_item, pursuit_title = await _pursuit_context_item(db, pursuit)
    items.append(context_item)

    organization = await db.get(Organization, organization_id)
    items.append(_manifest_item(
        category="FIRM", authority="Organization", identity=organization_id,
        provenance="ORGANIZATION_ACCOUNT", purpose="Approved scenario lead Organization",
        evidence_state="ACCOUNT_IDENTITY", payload={
            "role": "LEAD_ORGANIZATION",
            "display_name": organization.display_name if organization else None,
            "organization_id": organization_id,
        },
    ))

    pack_items = list((await db.scalars(select(AnalysisPackItem).where(
        AnalysisPackItem.pack_id == handoff.analysis_pack_id,
    ).order_by(AnalysisPackItem.ordinal))).all())
    pack_item_by_id = {item.id: item for item in pack_items}
    for source in pack_items:
        private = source.item_kind == "PRIVATE"
        items.append(_manifest_item(
            category="PRIVATE_DOCUMENT" if private else "SOURCE_DOCUMENT",
            authority="DocumentVersion" if private else "TenderDocument",
            identity=source.document_version_id if private else source.tender_document_id,
            provenance=source.provenance,
            purpose=f"Exact W4 input document: {source.display_name}",
            source_sha256=source.content_sha256,
            source_version=source.version_number,
            payload={
                "analysis_pack_item_id": source.id,
                "private_document_id": source.private_document_id,
                "document_version_id": source.document_version_id,
                "tender_document_id": source.tender_document_id,
                "display_name": source.display_name, "role": source.role,
                "source_url": source.source_url, "content_sha256": source.content_sha256,
                "identity_sha256": source.identity_sha256,
                "page_count": source.page_count, "page_count_known": source.page_count_known,
            },
        ))

    gap_by_requirement = {item.requirement_id: item for item in handoff.gap_assessments if item.requirement_id}
    gap_by_position = {item.position_id: item for item in handoff.gap_assessments if item.position_id}
    contributions_by_requirement: dict[UUID, list[tuple[Any, Any]]] = defaultdict(list)
    contributions_by_position: dict[UUID, list[tuple[Any, Any]]] = defaultdict(list)
    for participant in handoff.participants:
        for contribution in participant.contributions:
            if contribution.requirement_id:
                contributions_by_requirement[contribution.requirement_id].append((participant, contribution))
            if contribution.position_id:
                contributions_by_position[contribution.position_id].append((participant, contribution))

    matrix_rows = 0
    later_stage = 0
    checklist = 0
    requirement_terms = ("FORM", "ARTIFACT", "SIGNATURE", "BID_SECURITY", "BID SECURITY", "DOCUMENT")

    def add_matrix_item(target: Any, *, kind: str) -> None:
        nonlocal matrix_rows, later_stage, checklist
        is_requirement = kind == "REQUIREMENT"
        target_id = target.requirement_id if is_requirement else target.position_id
        gap = (gap_by_requirement if is_requirement else gap_by_position).get(target_id)
        contributions = (contributions_by_requirement if is_requirement else contributions_by_position).get(target_id, [])
        is_later = target.effective_coverage_state == "LATER_STAGE_OBLIGATION" or bool(gap and not gap.is_current_stage)
        label = target.effective_normalized_requirement if is_requirement else target.effective_title
        source = pack_item_by_id.get(target.pack_item_id)
        payload = {
            "matrix_kind": kind, "identity": target_id,
            "original_language_requirement": target.original_quote,
            "reviewed_label": label, "source_locator": target.source_locator,
            "source_pack_item_id": target.pack_item_id,
            "effective_review_state": target.effective_review_state,
            "effective_coverage_state": target.effective_coverage_state,
            "scenario_gap_outcome": gap.state if gap else (
                "COVERED" if target.effective_coverage_state in {"SUPPORTED", "NOT_APPLICABLE"} else "UNRESOLVED"
            ),
            "contributors": [{
                "participant_id": participant.participant_id,
                "candidate_id": participant.candidate_id,
                "participant_type": participant.participant_type,
                "display_name": participant.display_name,
                "candidate_match_id": contribution.candidate_match_id,
                "participation_record_id": contribution.participation_record_id,
                "supporting_evidence_identities": contribution.evidence_identities,
                "evidence_state": contribution.candidate_evidence_state,
                "participation_state": contribution.participation_effective_state,
            } for participant, contribution in contributions],
            "remaining_condition": None if (
                (gap and gap.state == "COVERED") or
                (not gap and target.effective_coverage_state in {"SUPPORTED", "NOT_APPLICABLE"})
            ) else (gap.rationale_code if gap else "No supporting evidence is sealed"),
            "later_stage": is_later,
        }
        if is_requirement:
            payload.update({
                "category": target.category, "requirement_type": target.requirement_type,
                "stage_scope": target.stage_scope, "distinction": target.distinction,
                "predicate": target.predicate,
            })
        else:
            payload.update({
                "quantity": target.quantity, "distinction": target.distinction,
                "education_qualification": target.education_qualification,
                "general_experience": target.general_experience,
                "specific_experience": target.specific_experience,
                "relevant_assignments": target.relevant_assignments,
                "languages": target.languages, "certifications": target.certifications,
                "expected_effort": target.expected_effort, "assignment_dates": target.assignment_dates,
            })
        items.append(_manifest_item(
            category="LATER_STAGE_OBLIGATION" if is_later else "REQUIREMENT",
            authority="PursuitRequirement" if is_requirement else "PursuitPosition",
            identity=target_id, provenance=source.provenance if source else "SEALED_W4",
            purpose=("Later-stage obligation" if is_later else "Requirement-to-evidence matrix row"),
            review_state=target.effective_review_state,
            evidence_state=target.effective_coverage_state,
            requirement_id=target_id if is_requirement else None,
            position_id=target_id if not is_requirement else None,
            gap_id=gap.gap_id if gap else None,
            source_sha256=source.content_sha256 if source else None,
            source_version=source.version_number if source else None,
            payload=payload,
        ))
        if is_later:
            later_stage += 1
        else:
            matrix_rows += 1
        searchable = " ".join(str(value or "") for value in (
            label, getattr(target, "category", None), getattr(target, "requirement_type", None),
        )).upper()
        if is_requirement and any(term in searchable for term in requirement_terms):
            checklist_state = "AVAILABLE" if source and source.role == "FORM" else "NEEDS_ACTION"
            items.append(_manifest_item(
                category="FORM_OR_REQUIRED_ARTIFACT", authority="PursuitRequirement", identity=target_id,
                provenance=source.provenance if source else "SEALED_W4",
                purpose="Source-backed required form or artifact checklist item",
                review_state=target.effective_review_state, evidence_state=checklist_state,
                requirement_id=target_id, gap_id=gap.gap_id if gap else None,
                source_sha256=source.content_sha256 if source else None,
                payload={
                    "label": label, "checklist_state": checklist_state,
                    "source_pack_item_id": target.pack_item_id,
                    "linked_document_identity": str(
                        source.document_version_id if source and source.item_kind == "PRIVATE"
                        else source.tender_document_id if source else ""
                    ) or None,
                    "automatically_completed": False,
                },
            ))
            checklist += 1

    for requirement in analysis.requirements:
        add_matrix_item(requirement, kind="REQUIREMENT")
    for position in analysis.positions:
        add_matrix_item(position, kind="POSITION")

    for assessment in handoff.gap_assessments:
        items.append(_manifest_item(
            category="GAP_ASSESSMENT", authority="ScenarioGapAssessment",
            identity=assessment.gap_assessment_id, provenance="SEALED_W7",
            purpose="Exact W7 scenario outcome for the reviewed W4 Gap",
            review_state=assessment.w4_review_state, evidence_state=assessment.state,
            requirement_id=assessment.requirement_id, position_id=assessment.position_id,
            gap_id=assessment.gap_id, payload=assessment.model_dump(mode="json"),
        ))

    reference_ids, cv_ids = _selected_evidence_ids(handoff)
    references = list((await db.scalars(select(ProjectReference).where(
        ProjectReference.id.in_(reference_ids)
    ).order_by(ProjectReference.id))).all()) if reference_ids else []
    cvs = list((await db.scalars(select(CVVersion).where(
        CVVersion.id.in_(cv_ids)
    ).order_by(CVVersion.id))).all()) if cv_ids else []
    reference_by_id = {item.id: item for item in references}
    cv_by_id = {item.id: item for item in cvs}

    for participant in handoff.participants:
        participant_category = "FIRM" if participant.participant_type == "PARTNER_FIRM" else "EXPERT"
        items.append(_manifest_item(
            category=participant_category,
            authority="Firm" if participant.participant_type == "PARTNER_FIRM" else "Expert",
            identity=participant.candidate_id, provenance="SEALED_W7_SELECTION",
            purpose="Approved scenario Partner Firm" if participant.participant_type == "PARTNER_FIRM" else "Approved scenario Expert",
            payload={
                "participant_id": participant.participant_id,
                "candidate_id": participant.candidate_id,
                "display_name": participant.display_name,
                "participant_type": participant.participant_type,
                "contributions": [{
                    "contribution_id": value.contribution_id,
                    "gap_id": value.gap_id,
                    "requirement_id": value.requirement_id,
                    "position_id": value.position_id,
                    "proposed_contribution": value.proposed_contribution,
                    "qualification_state": value.qualification_state,
                    "evidence_state": value.candidate_evidence_state,
                    "evidence_identities": value.evidence_identities,
                } for value in participant.contributions],
            },
        ))
        seen_refs: set[UUID] = set()
        seen_cvs: set[UUID] = set()
        for contribution in participant.contributions:
            items.append(_manifest_item(
                category="PARTICIPATION_CONFIRMATION", authority="CandidateParticipationRecord",
                identity=contribution.participation_record_id, provenance="SEALED_W6_VIA_W7",
                purpose="Exact participation facts authorizing this proposed contribution",
                review_state=contribution.shortlist_decision_state,
                evidence_state=contribution.participation_effective_state,
                requirement_id=contribution.requirement_id, position_id=contribution.position_id,
                gap_id=contribution.gap_id,
                payload={
                    "candidate_match_id": contribution.candidate_match_id,
                    "participation_record_id": contribution.participation_record_id,
                    "availability_fact_id": contribution.availability_fact_id,
                    "availability_recorded_state": contribution.availability_recorded_state,
                    "availability_effective_state": contribution.availability_effective_state,
                    "availability_window_start": contribution.availability_window_start,
                    "availability_window_end": contribution.availability_window_end,
                    "availability_effort_percent": contribution.availability_effort_percent,
                    "availability_capacity": contribution.availability_capacity,
                    "availability_valid_until": contribution.availability_valid_until,
                    "interest_fact_id": contribution.interest_fact_id,
                    "interest_recorded_state": contribution.interest_recorded_state,
                    "interest_effective_state": contribution.interest_effective_state,
                    "interest_conditions": contribution.interest_conditions,
                    "participation_decision_id": contribution.participation_decision_id,
                    "participation_recorded_state": contribution.participation_recorded_state,
                    "participation_effective_state": contribution.participation_effective_state,
                    "confirmation_source": contribution.confirmation_source,
                    "confirmation_observed_at": contribution.confirmation_observed_at,
                    "reconfirm_by": contribution.reconfirm_by,
                    "confirmation_conditions": contribution.confirmation_conditions,
                    "assignment_dates": contribution.assignment_dates,
                    "assignment_window_result": contribution.assignment_window_result,
                    # W7 seals the classification claimed by the recorder, but it
                    # does not seal the supporting document identity. Do not turn
                    # a SIGNED_DOCUMENT label into an independent verification.
                    "provenance_independently_verified": False,
                    "supporting_document_identity_sealed": False,
                },
            ))
            for identity in contribution.evidence_identities:
                try:
                    evidence_id = UUID(str(identity.get("id")))
                except (ValueError, TypeError, AttributeError):
                    continue
                if identity.get("type") == "PROJECT_REFERENCE" and evidence_id not in seen_refs:
                    seen_refs.add(evidence_id)
                    reference = reference_by_id.get(evidence_id)
                    if reference:
                        items.append(_manifest_item(
                            category="PROJECT_REFERENCE", authority="ProjectReference",
                            identity=reference.id, provenance="REVIEWED_STRUCTURED_REFERENCE",
                            purpose="Selected Firm reference supporting the exact contribution",
                            evidence_state=reference.evidence_state,
                            requirement_id=contribution.requirement_id, gap_id=contribution.gap_id,
                            payload={
                                "firm_id": reference.firm_id, "project": reference.project_name,
                                "client": reference.client_name, "role": reference.role,
                                "firm_share_percent": reference.contract_share_percent,
                                "value": reference.contract_value, "currency": reference.contract_currency,
                                "value_basis": reference.value_basis, "start_date": reference.start_date,
                                "completion_date": reference.completion_date,
                                "completion_state": reference.completion_state,
                                "relevant_scope": reference.relevant_scope,
                                "evidence_state": reference.evidence_state,
                                "evidence_provenance": reference.evidence_provenance,
                                "attached_contract": False,
                                "value_is_firm_share": reference.value_basis == "FIRM_SHARE",
                            },
                        ))
                elif identity.get("type") == "CV_VERSION" and evidence_id not in seen_cvs:
                    seen_cvs.add(evidence_id)
                    cv = cv_by_id.get(evidence_id)
                    if cv:
                        items.append(_manifest_item(
                            category="CV_FACTS", authority="CVVersion", identity=cv.id,
                            provenance="STRUCTURED_CV_FACTS",
                            purpose="Structured expert fact sheet / proposal preparation draft",
                            evidence_state=cv.evidence_state,
                            position_id=contribution.position_id, gap_id=contribution.gap_id,
                            source_sha256=cv.structured_sha256, source_version=cv.version_number,
                            payload={
                                "expert_id": cv.expert_id, "cv_version_id": cv.id,
                                "version_number": cv.version_number,
                                "structured_sha256": cv.structured_sha256,
                                "education": cv.education, "qualifications": cv.qualifications,
                                "certifications": cv.certifications, "assignments": cv.assignments,
                                "languages": cv.languages, "evidence_state": cv.evidence_state,
                                "evidence_provenance": cv.evidence_provenance,
                                "document_classification": "STRUCTURED_CV_FACTS",
                                "source_cv_document_id": None,
                                "source_cv_document_available": False,
                                "candidate_signed": False, "original_cv": False,
                            },
                        ))

    company = await db.scalar(select(AnalysisCompanySnapshot).where(
        AnalysisCompanySnapshot.analysis_run_id == handoff.analysis_run_id
    ))
    if company is not None:
        items.append(_manifest_item(
            category="OTHER", authority="AnalysisCompanySnapshot", identity=company.id,
            provenance="SEALED_W4_COMPANY_SNAPSHOT",
            purpose="Company/readiness evidence with its accepted authority basis preserved",
            evidence_state="MIXED_AUTHORITY",
            source_sha256=company.snapshot_sha256,
            payload={
                "legacy_company_profile_id": company.legacy_company_profile_id,
                "snapshot": company.snapshot_json,
                "metadata_only_count": company.metadata_only_count,
                "file_backed_count": company.file_backed_count,
                "metadata_promoted_to_verified": False,
            },
        ))

    return items, pursuit_title, {
        "matrix": matrix_rows, "participants": len(handoff.participants),
        "later": later_stage, "checklist": checklist,
    }


async def seal_proposal_evidence_pack(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID,
    membership_id: UUID, request: ProposalEvidenceSealRequest,
) -> ProposalEvidencePackResponse:
    await _lock_sealing_authorities(db)
    await _active_membership(db, organization_id, membership_id)
    scenario = await db.scalar(select(TeamScenario).where(
        TeamScenario.id == request.scenario_id,
        TeamScenario.organization_id == organization_id,
        TeamScenario.pursuit_id == pursuit_id,
    ))
    if scenario is None:
        raise ProposalEvidenceNotFoundError("Approved Team Scenario not found")
    try:
        handoff = await get_proposal_handoff(
            db, organization_id=organization_id, pursuit_id=pursuit_id,
            scenario_id=request.scenario_id, revision_id=request.revision_id,
        )
    except TeamScenarioEligibilityError as exc:
        raise ProposalEvidenceEligibilityError(str(exc)) from exc
    if handoff.approval_decision_id != request.approval_decision_id:
        raise ProposalEvidenceEligibilityError("The exact current approval decision does not match")
    if not handoff.scenario_current or handoff.assessment_state != "VIABLE":
        raise ProposalEvidenceEligibilityError("Sealing requires a current VIABLE W7 handoff")

    manifest, pursuit_title, counts = await _build_manifest(
        db, organization_id=organization_id, pursuit_id=pursuit_id,
        scenario=scenario, handoff=handoff,
    )
    canonical_manifest = {
        "schema_version": PACK_SCHEMA_VERSION,
        "organization_id": organization_id, "pursuit_id": pursuit_id,
        "scenario_id": request.scenario_id, "scenario_revision_id": request.revision_id,
        "approval_decision_id": request.approval_decision_id,
        "analysis_run_id": handoff.analysis_run_id, "analysis_pack_id": handoff.analysis_pack_id,
        "items": manifest,
    }
    manifest_sha = hashlib.sha256(_canonical(canonical_manifest)).hexdigest()
    workspace = await db.scalar(select(PursuitProposalWorkspace).where(
        PursuitProposalWorkspace.organization_id == organization_id,
        PursuitProposalWorkspace.pursuit_id == pursuit_id,
    ))
    if workspace is None:
        workspace = PursuitProposalWorkspace(
            organization_id=organization_id, pursuit_id=pursuit_id,
            created_by_membership_id=membership_id,
        )
        db.add(workspace)
        await db.flush()
    else:
        workspace.updated_at = datetime.now(timezone.utc)
    version = (await db.scalar(select(func.max(ProposalEvidencePack.pack_version)).where(
        ProposalEvidencePack.workspace_id == workspace.id
    ))) or 0
    pack = ProposalEvidencePack(
        organization_id=organization_id, pursuit_id=pursuit_id, workspace_id=workspace.id,
        pack_version=version + 1, scenario_id=request.scenario_id,
        scenario_revision_id=request.revision_id,
        approval_decision_id=request.approval_decision_id,
        analysis_run_id=handoff.analysis_run_id, analysis_pack_id=handoff.analysis_pack_id,
        sealed_by_membership_id=membership_id, schema_version=PACK_SCHEMA_VERSION,
        manifest_sha256=manifest_sha, pack_state="SEALED",
        scenario_title_snapshot=scenario.title, pursuit_title_snapshot=pursuit_title,
        item_count=len(manifest), matrix_row_count=counts["matrix"],
        participant_count=counts["participants"], later_stage_count=counts["later"],
        checklist_count=counts["checklist"],
    )
    db.add(pack)
    await db.flush()
    rows = [ProposalEvidencePackItem(
        organization_id=organization_id, pack_id=pack.id, ordinal=index,
        category=item["category"], source_authority_type=item["source_authority_type"],
        source_identity=item["source_identity"], provenance=item["provenance"],
        review_state=item["review_state"], evidence_state=item["evidence_state"],
        purpose=item["purpose"], requirement_id=item["requirement_id"],
        position_id=item["position_id"], gap_id=item["gap_id"],
        source_sha256=item["source_sha256"], source_version=item["source_version"],
        payload_snapshot=item["payload"],
    ) for index, item in enumerate(manifest, start=1)]
    db.add_all(rows)
    await db.flush()

    # Revalidate once more while all append authorities remain locked. The browser never supplies truth.
    try:
        final_handoff = await get_proposal_handoff(
            db, organization_id=organization_id, pursuit_id=pursuit_id,
            scenario_id=request.scenario_id, revision_id=request.revision_id,
        )
    except TeamScenarioEligibilityError as exc:
        await db.rollback()
        raise ProposalEvidenceEligibilityError("The W7 handoff changed while sealing") from exc
    if final_handoff.approval_decision_id != request.approval_decision_id:
        await db.rollback()
        raise ProposalEvidenceEligibilityError("The approval decision changed while sealing")
    await db.commit()
    result = await get_proposal_evidence_pack(
        db, organization_id=organization_id, pursuit_id=pursuit_id, pack_id=pack.id,
    )
    if result is None:
        raise RuntimeError("Sealed Proposal Evidence Pack could not be projected")
    return result


def _item_response(item: ProposalEvidencePackItem) -> ProposalEvidenceItemResponse:
    return ProposalEvidenceItemResponse(
        item_id=item.id, ordinal=item.ordinal, category=item.category,
        source_authority_type=item.source_authority_type,
        source_identity=item.source_identity, provenance=item.provenance,
        review_state=item.review_state, evidence_state=item.evidence_state,
        purpose=item.purpose, requirement_id=item.requirement_id,
        position_id=item.position_id, gap_id=item.gap_id,
        source_sha256=item.source_sha256, source_version=item.source_version,
        payload=item.payload_snapshot,
    )


def _artifact_response(item: ProposalEvidenceArtifact) -> ProposalEvidenceArtifactResponse:
    return ProposalEvidenceArtifactResponse(
        artifact_id=item.id, pack_id=item.pack_id, artifact_type=item.artifact_type,
        historical_snapshot=item.historical_snapshot,
        content_sha256=item.content_sha256, byte_size=item.byte_size,
        media_type=item.media_type, generator_version=item.generator_version,
        created_by_membership_id=item.created_by_membership_id,
        created_at=item.created_at,
    )


async def _project_packs(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID,
    packs: list[ProposalEvidencePack],
) -> list[ProposalEvidencePackResponse]:
    if not packs:
        return []
    pack_ids = [item.id for item in packs]
    items = list((await db.scalars(select(ProposalEvidencePackItem).where(
        ProposalEvidencePackItem.pack_id.in_(pack_ids),
        ProposalEvidencePackItem.organization_id == organization_id,
    ).order_by(ProposalEvidencePackItem.pack_id, ProposalEvidencePackItem.ordinal))).all())
    artifacts = list((await db.scalars(select(ProposalEvidenceArtifact).where(
        ProposalEvidenceArtifact.pack_id.in_(pack_ids),
        ProposalEvidenceArtifact.organization_id == organization_id,
    ).order_by(ProposalEvidenceArtifact.pack_id, ProposalEvidenceArtifact.created_at, ProposalEvidenceArtifact.id))).all())
    items_by_pack: dict[UUID, list[ProposalEvidencePackItem]] = defaultdict(list)
    artifacts_by_pack: dict[UUID, list[ProposalEvidenceArtifact]] = defaultdict(list)
    for item in items:
        items_by_pack[item.pack_id].append(item)
    for artifact in artifacts:
        artifacts_by_pack[artifact.pack_id].append(artifact)

    scenario_ids = sorted({item.scenario_id for item in packs}, key=str)
    scenarios = list((await db.scalars(select(TeamScenario).where(
        TeamScenario.id.in_(scenario_ids),
        TeamScenario.organization_id == organization_id,
        TeamScenario.pursuit_id == pursuit_id,
    ).order_by(TeamScenario.id))).all())
    projected_scenarios = await _project_scenarios(
        db, organization_id=organization_id, pursuit_id=pursuit_id, scenarios=scenarios,
    ) if scenarios else []
    scenario_by_id = {item.scenario_id: item for item in projected_scenarios}

    result: list[ProposalEvidencePackResponse] = []
    for pack in packs:
        reasons: list[str] = []
        scenario = scenario_by_id.get(pack.scenario_id)
        revision = next((
            item for item in (scenario.revisions if scenario else [])
            if item.revision_id == pack.scenario_revision_id
        ), None)
        if revision is None:
            reasons.append("The exact W7 scenario revision is unavailable.")
        else:
            if not revision.scenario_current:
                reasons.extend(revision.stale_reasons or ["The W7 scenario revision is stale."])
            if revision.current_assessment_state != "VIABLE":
                reasons.append("The exact W7 scenario revision is no longer VIABLE.")
            latest_decision = revision.decisions[-1] if revision.decisions else None
            if (
                latest_decision is None
                or latest_decision.decision != "APPROVED_FOR_PROPOSAL"
                or latest_decision.decision_id != pack.approval_decision_id
            ):
                reasons.append("The exact W7 approval is no longer the current revision decision.")
        result.append(ProposalEvidencePackResponse(
            pack_id=pack.id, workspace_id=pack.workspace_id,
            organization_id=pack.organization_id, pursuit_id=pack.pursuit_id,
            pack_version=pack.pack_version, scenario_id=pack.scenario_id,
            scenario_revision_id=pack.scenario_revision_id,
            approval_decision_id=pack.approval_decision_id,
            analysis_run_id=pack.analysis_run_id, analysis_pack_id=pack.analysis_pack_id,
            schema_version=pack.schema_version, manifest_sha256=pack.manifest_sha256,
            pack_state=pack.pack_state, scenario_title=pack.scenario_title_snapshot,
            pursuit_title=pack.pursuit_title_snapshot, item_count=pack.item_count,
            matrix_row_count=pack.matrix_row_count, participant_count=pack.participant_count,
            later_stage_count=pack.later_stage_count, checklist_count=pack.checklist_count,
            pack_current=not reasons, stale_reasons=list(dict.fromkeys(reasons)),
            items=[_item_response(item) for item in items_by_pack[pack.id]],
            artifacts=[_artifact_response(item) for item in artifacts_by_pack[pack.id]],
            sealed_by_membership_id=pack.sealed_by_membership_id,
            created_at=pack.created_at,
        ))
    return result


async def get_proposal_workspace(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID, user_id: UUID,
) -> PursuitProposalWorkspaceResponse:
    pursuit = await db.scalar(select(OrganizationPursuit).where(
        OrganizationPursuit.id == pursuit_id,
        OrganizationPursuit.organization_id == organization_id,
    ))
    if pursuit is None:
        raise ProposalEvidenceNotFoundError("Pursuit not found")
    workspace = await db.scalar(select(PursuitProposalWorkspace).where(
        PursuitProposalWorkspace.organization_id == organization_id,
        PursuitProposalWorkspace.pursuit_id == pursuit_id,
    ))
    packs = list((await db.scalars(select(ProposalEvidencePack).where(
        ProposalEvidencePack.organization_id == organization_id,
        ProposalEvidencePack.pursuit_id == pursuit_id,
    ).order_by(ProposalEvidencePack.pack_version.desc()).limit(100))).all())
    projected = await _project_packs(
        db, organization_id=organization_id, pursuit_id=pursuit_id, packs=packs,
    )
    legacy_id = None
    if pursuit.origin == PursuitOrigin.SOURCE and pursuit.source_tender_id:
        legacy_id = await db.scalar(select(Proposal.id).where(
            Proposal.user_id == user_id, Proposal.tender_id == pursuit.source_tender_id,
        ).limit(1))
    return PursuitProposalWorkspaceResponse(
        workspace_id=workspace.id if workspace else None,
        organization_id=organization_id, pursuit_id=pursuit_id,
        created_by_membership_id=workspace.created_by_membership_id if workspace else None,
        created_at=workspace.created_at if workspace else None,
        updated_at=workspace.updated_at if workspace else None,
        legacy_proposal_id=legacy_id, packs=projected,
    )


async def get_proposal_evidence_pack(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID, pack_id: UUID,
) -> ProposalEvidencePackResponse | None:
    pack = await db.scalar(select(ProposalEvidencePack).where(
        ProposalEvidencePack.id == pack_id,
        ProposalEvidencePack.organization_id == organization_id,
        ProposalEvidencePack.pursuit_id == pursuit_id,
    ))
    if pack is None:
        return None
    return (await _project_packs(
        db, organization_id=organization_id, pursuit_id=pursuit_id, packs=[pack],
    ))[0]


def _export_manifest(pack: ProposalEvidencePackResponse, historical: bool) -> dict[str, Any]:
    return {
        "document_classification": "HISTORICAL SNAPSHOT" if historical else "CURRENT SEALED SNAPSHOT",
        "disclosure": {
            "source_text": "Quoted text is retained from exact sealed W4 source inputs.",
            "reviewed_structured_fact": "Structured facts retain their source authority and review state.",
            "plasma_generated_summary": "Headings and arrangement are deterministic summaries of sealed pack data.",
            "commercial_price_included": False,
            "compliance_certificate": False,
        },
        "pack": {
            "pack_id": pack.pack_id, "pack_version": pack.pack_version,
            "manifest_sha256": pack.manifest_sha256, "schema_version": pack.schema_version,
            "created_at": pack.created_at, "pursuit_id": pack.pursuit_id,
            "pursuit_title": pack.pursuit_title, "scenario_id": pack.scenario_id,
            "scenario_revision_id": pack.scenario_revision_id,
            "approval_decision_id": pack.approval_decision_id,
            "analysis_run_id": pack.analysis_run_id, "analysis_pack_id": pack.analysis_pack_id,
        },
        "items": [{
            "ordinal": item.ordinal, "category": item.category,
            "source_authority_type": item.source_authority_type,
            "source_identity": item.source_identity, "provenance": item.provenance,
            "review_state": item.review_state, "evidence_state": item.evidence_state,
            "purpose": item.purpose, "requirement_id": item.requirement_id,
            "position_id": item.position_id, "gap_id": item.gap_id,
            "source_sha256": item.source_sha256, "source_version": item.source_version,
            "payload": item.payload,
        } for item in pack.items],
    }


def _sections(manifest: dict[str, Any]) -> list[tuple[str, list[dict[str, Any]]]]:
    labels = [
        ("PURSUIT_CONTEXT", "Pursuit context"),
        ("REQUIREMENT", "Requirement-to-evidence matrix"),
        ("FIRM", "Lead Organization and Partner Firms"),
        ("EXPERT", "Experts"),
        ("PROJECT_REFERENCE", "Firm reference evidence"),
        ("CV_FACTS", "Structured expert facts"),
        ("PARTICIPATION_CONFIRMATION", "Participation confirmation"),
        ("SOURCE_DOCUMENT", "Source document manifest"),
        ("PRIVATE_DOCUMENT", "Private document manifest"),
        ("LATER_STAGE_OBLIGATION", "Later-stage obligations"),
        ("FORM_OR_REQUIRED_ARTIFACT", "Required forms and artifacts"),
        ("GAP_ASSESSMENT", "Gap assessments"),
        ("OTHER", "Other reviewed evidence"),
    ]
    return [(label, [item for item in manifest["items"] if item["category"] == category]) for category, label in labels]


class _InvariantCanvas(Canvas):
    def __init__(self, *args: Any, **kwargs: Any) -> None:
        kwargs["invariant"] = 1
        super().__init__(*args, **kwargs)


def _pdf_bytes(manifest: dict[str, Any]) -> bytes:
    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title="Proposal Evidence Summary", author="Plasma",
    )
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph("Proposal Evidence Summary", styles["Title"]),
        Paragraph(escape(manifest["document_classification"]), styles["Heading2"]),
        Paragraph(escape(str(manifest["pack"]["pursuit_title"])), styles["Heading2"]),
        Paragraph(f"Pack {manifest['pack']['pack_version']} · {escape(str(manifest['pack']['manifest_sha256']))}", styles["BodyText"]),
        Spacer(1, 8),
    ]
    for title, rows in _sections(manifest):
        if not rows:
            continue
        story.append(Paragraph(escape(title), styles["Heading2"]))
        table_rows = [["Authority", "Purpose", "State"]]
        for item in rows:
            label = item["payload"].get("reviewed_label") or item["payload"].get("display_name") or item["source_identity"]
            table_rows.append([
                Paragraph(escape(f"{item['source_authority_type']}: {label}"), styles["BodyText"]),
                Paragraph(escape(str(item["purpose"])), styles["BodyText"]),
                Paragraph(escape(str(item["evidence_state"] or item["review_state"] or "DISCLOSED")), styles["BodyText"]),
            ])
        table = Table(table_rows, colWidths=[58 * mm, 80 * mm, 30 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#A9B4C6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([table, Spacer(1, 8)])
    disclosure = manifest["disclosure"]
    story.extend([
        Paragraph("Provenance and disclosure", styles["Heading2"]),
        Paragraph(escape(disclosure["source_text"]), styles["BodyText"]),
        Paragraph(escape(disclosure["reviewed_structured_fact"]), styles["BodyText"]),
        Paragraph(escape(disclosure["plasma_generated_summary"]), styles["BodyText"]),
        Paragraph("Commercial price included: No", styles["BodyText"]),
    ])
    document.build(story, canvasmaker=_InvariantCanvas)
    return stream.getvalue()


def _normalize_docx(raw: bytes) -> bytes:
    source = ZipFile(io.BytesIO(raw), "r")
    target_stream = io.BytesIO()
    with source, ZipFile(target_stream, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return target_stream.getvalue()


def _docx_bytes(manifest: dict[str, Any]) -> bytes:
    document = Document()
    document.core_properties.title = "Proposal Evidence Preparation Brief"
    document.core_properties.author = "Plasma"
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    document.add_heading("Proposal Evidence Preparation Brief", 0)
    document.add_paragraph(manifest["document_classification"])
    document.add_heading(str(manifest["pack"]["pursuit_title"]), level=1)
    document.add_paragraph(
        f"Pack {manifest['pack']['pack_version']} · Manifest {manifest['pack']['manifest_sha256']}"
    )
    for title, rows in _sections(manifest):
        if not rows:
            continue
        document.add_heading(title, level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for index, value in enumerate(("Authority", "Purpose", "State")):
            table.rows[0].cells[index].text = value
        for item in rows:
            cells = table.add_row().cells
            label = item["payload"].get("reviewed_label") or item["payload"].get("display_name") or item["source_identity"]
            cells[0].text = f"{item['source_authority_type']}: {label}"
            cells[1].text = str(item["purpose"])
            cells[2].text = str(item["evidence_state"] or item["review_state"] or "DISCLOSED")
    document.add_heading("Provenance and disclosure", level=1)
    for value in manifest["disclosure"].values():
        if isinstance(value, str):
            document.add_paragraph(value)
    document.add_paragraph("Commercial price included: No")
    stream = io.BytesIO()
    document.save(stream)
    return _normalize_docx(stream.getvalue())


def _artifact_content(pack: ProposalEvidencePackResponse, artifact_type: str, historical: bool) -> tuple[bytes, str]:
    manifest = _export_manifest(pack, historical)
    if artifact_type == "JSON":
        return _canonical(manifest) + b"\n", "application/json"
    if artifact_type == "PDF":
        return _pdf_bytes(manifest), "application/pdf"
    return _docx_bytes(manifest), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _artifact_storage_key(
    *, organization_id: UUID, pursuit_id: UUID, pack_id: UUID, artifact_id: UUID,
) -> str:
    return f"exports/{organization_id.hex}/{pursuit_id.hex}/{pack_id.hex}/{artifact_id.hex}.bin"


def _write_private_artifact(storage_key: str, content: bytes) -> Path:
    target = resolve_private_storage_key(storage_key)
    target.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    if target.exists():
        raise ProposalEvidenceEligibilityError("Private artifact storage identity already exists")
    staged: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix="export-", suffix=".tmp", delete=False) as stream:
            staged = Path(stream.name)
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(staged, target)
        target.chmod(0o600)
        return target
    except Exception:
        if staged is not None:
            staged.unlink(missing_ok=True)
        raise


async def generate_proposal_evidence_artifact(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID, pack_id: UUID,
    membership_id: UUID, request: ProposalEvidenceExportRequest,
) -> ProposalEvidenceArtifactResponse:
    await _active_membership(db, organization_id, membership_id)
    pack = await get_proposal_evidence_pack(
        db, organization_id=organization_id, pursuit_id=pursuit_id, pack_id=pack_id,
    )
    if pack is None:
        raise ProposalEvidenceNotFoundError("Proposal Evidence Pack not found")
    if not pack.pack_current and not request.historical_snapshot:
        raise ProposalEvidenceEligibilityError(
            "A stale pack can create only an explicitly labeled HISTORICAL SNAPSHOT export"
        )
    existing = await db.scalar(select(ProposalEvidenceArtifact).where(
        ProposalEvidenceArtifact.organization_id == organization_id,
        ProposalEvidenceArtifact.pack_id == pack_id,
        ProposalEvidenceArtifact.artifact_type == request.artifact_type,
        ProposalEvidenceArtifact.historical_snapshot == request.historical_snapshot,
    ))
    if existing is not None:
        return _artifact_response(existing)
    content, media_type = _artifact_content(pack, request.artifact_type, request.historical_snapshot)
    artifact_id = uuid4()
    storage_key = _artifact_storage_key(
        organization_id=organization_id, pursuit_id=pursuit_id,
        pack_id=pack_id, artifact_id=artifact_id,
    )
    target: Path | None = None
    try:
        target = _write_private_artifact(storage_key, content)
        artifact = ProposalEvidenceArtifact(
            id=artifact_id, organization_id=organization_id, pack_id=pack_id,
            artifact_type=request.artifact_type,
            historical_snapshot=request.historical_snapshot,
            storage_key=storage_key, content_sha256=hashlib.sha256(content).hexdigest(),
            byte_size=len(content), media_type=media_type,
            generator_version=EXPORT_GENERATOR_VERSION,
            created_by_membership_id=membership_id,
        )
        db.add(artifact)
        await db.commit()
        await db.refresh(artifact)
        return _artifact_response(artifact)
    except Exception:
        await db.rollback()
        if target is not None:
            target.unlink(missing_ok=True)
        raise


async def resolve_proposal_evidence_artifact(
    db: AsyncSession, *, organization_id: UUID, pursuit_id: UUID,
    artifact_id: UUID, membership_id: UUID,
) -> tuple[ProposalEvidenceArtifact, Path]:
    await _active_membership(db, organization_id, membership_id)
    artifact = await db.scalar(select(ProposalEvidenceArtifact).join(
        ProposalEvidencePack, ProposalEvidencePack.id == ProposalEvidenceArtifact.pack_id,
    ).where(
        ProposalEvidenceArtifact.id == artifact_id,
        ProposalEvidenceArtifact.organization_id == organization_id,
        ProposalEvidencePack.pursuit_id == pursuit_id,
        ProposalEvidencePack.organization_id == organization_id,
    ))
    if artifact is None:
        raise ProposalEvidenceNotFoundError("Proposal Evidence artifact not found")
    path = resolve_private_storage_key(artifact.storage_key)
    if not path.is_file():
        raise ProposalEvidenceNotFoundError("Proposal Evidence artifact bytes are unavailable")
    if path.stat().st_size != artifact.byte_size:
        raise ProposalEvidenceEligibilityError("Proposal Evidence artifact integrity check failed")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != artifact.content_sha256:
        raise ProposalEvidenceEligibilityError("Proposal Evidence artifact integrity check failed")
    return artifact, path

