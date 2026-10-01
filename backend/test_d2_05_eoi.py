"""D2-05 Expression of Interest: contract, passivity, validation, isolation, immutability,
staleness, rendering (en/ru, DOCX/PDF) and the relevance-note validator and budget."""

from __future__ import annotations

import asyncio
from datetime import date
from decimal import Decimal
import hashlib
import io
from pathlib import Path
import re
from uuid import uuid4

import fitz
from docx import Document
import pytest
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.agents.pursuit_analyzer import ExtractedFact, VerifiedFact
from app.core.config import settings
from app.main import app
from app.models.all_models import TenderDocument
from app.models.base import TenderEngagementOrigin, TenderEngagementStatus
from app.models.eoi import EoiDraft, EoiDraftArtifact
from app.models.pursuit_analysis import AnalysisRun
from app.schemas.candidate_retrieval import (
    FirmCreateRequest,
    ProjectReferenceCreateRequest,
    ProjectReferenceUpdateRequest,
    SelfFirmUpsertRequest,
)
from app.schemas.eoi import (
    EoiDraftCreateRequest,
    EoiDraftResponse,
    EoiReference,
    EoiSuggestionsResponse,
)
from app.schemas.tenancy import PursuitAnalysisStartRequest
from app.services import eoi as eoi_service
from app.services import eoi_notes
from app.services import pursuit_analysis as analysis_service
from app.services.candidate_retrieval import (
    archive_project_reference,
    create_firm,
    create_project_reference,
    update_project_reference,
    upsert_self_firm,
)
from app.services.eoi import (
    EoiConflictError,
    EoiNotFoundError,
    create_eoi_draft,
    eoi_suggestions,
    get_eoi_draft,
    list_eoi_drafts,
    resolve_eoi_artifact,
)
from app.services.eoi_document import docx_bytes, document_text, pdf_bytes
from app.services.eoi_notes import NoteRequest, generate_relevance_notes, validate_note
from app.services.private_documents import build_analysis_pack_candidate
from app.services.pursuit_analysis import create_analysis_run, process_analysis_run
from app.services.pursuits import get_or_create_source_pursuit
from scripts import test_s0_5b4_baseline as support
from test_d2_01_own_experience import SUBSTATION, _fact
from test_w2_organization_pursuit_foundation import W1_HEAD, _seed_w1


D2_01_HEAD = "20261004_0001_d2_01_own_experience"
HEAD = "20261005_0001_d2_05_eoi_drafts"
PRICE_WORDS = re.compile(r"\bprice|\bpricing|\bfee\b|\bfees\b|remuneration|financial proposal|цен[аыу]|стоимость услуг|вознагражд", re.I)


# ---- contract ---------------------------------------------------------------------------------------

def test_contract_shapes_match_the_frontend_contract() -> None:
    paths = app.openapi()["paths"]
    assert set(paths["/api/v1/pursuits/{pursuit_id}/eoi/suggestions"]) == {"get"}
    assert set(paths["/api/v1/pursuits/{pursuit_id}/eoi-drafts"]) == {"get", "post"}
    assert "201" in paths["/api/v1/pursuits/{pursuit_id}/eoi-drafts"]["post"]["responses"]
    assert set(paths["/api/v1/pursuits/{pursuit_id}/eoi-drafts/{draft_id}"]) == {"get"}
    assert set(paths["/api/v1/pursuits/{pursuit_id}/eoi-artifacts/{artifact_id}/download"]) == {"get"}
    assert set(EoiSuggestionsResponse.model_fields) == {
        "analysis_run_id", "run_current", "defaults", "criteria", "notes", "own_references", "partner_firms",
    }
    assert set(EoiReference.model_fields) == {
        "reference_id", "project_name", "client_name", "country", "sector", "service", "role",
        "contract_share_percent", "contract_value", "contract_currency", "value_basis", "start_date",
        "completion_date", "completion_state", "relevant_scope", "evidence_state", "evidence_basis",
        "matched_requirement_ids", "suggested", "rank",
    }
    assert set(EoiDraftResponse.model_fields) == {
        "draft_id", "version", "created_at", "created_by_membership_id", "analysis_run_id", "language",
        "current", "stale_reasons", "artifacts", "summary",
    }
    schemas = app.openapi()["components"]["schemas"]
    assert set(schemas["EoiDefaults"]["properties"]) == {
        "assignment_title", "reference_no", "addressee_organization", "addressee_name", "addressee_email",
        "firm_name", "firm_country",
    }
    assert set(schemas["EoiDraftSummary"]["properties"]) == {
        "criteria_total", "criteria_with_references", "criteria_without_references", "own_reference_count",
        "partner_count", "relevance_notes_generated", "relevance_notes_dropped",
    }
    letter = {"addressee_organization": "Ministry", "signatory_name": "A. Signer", "signatory_title": "Director",
              "contact_email": "a@b.uz"}
    base = {"analysis_run_id": str(uuid4()), "language": "en", "own_reference_ids": [str(uuid4())], "letter": letter}
    assert EoiDraftCreateRequest(**base).include_relevance_notes is True
    one = str(uuid4())
    for bad in (
        {**base, "own_reference_ids": []},
        {**base, "own_reference_ids": [one, one]},
        {**base, "own_reference_ids": [str(uuid4()) for _ in range(31)]},
        {**base, "language": "uz"},
        {**base, "partners": [{"firm_id": str(uuid4()), "role": "LEAD", "reference_ids": []}]},
        {**base, "partners": [{"firm_id": str(uuid4()), "role": "JV_MEMBER", "reference_ids": []} for _ in range(6)]},
        {**base, "partners": [{"firm_id": one, "role": "JV_MEMBER"}, {"firm_id": one, "role": "SUBCONSULTANT"}]},
        {**base, "letter": {**letter, "contact_email": "not-an-email"}},
        {**base, "letter": {**letter, "signatory_name": "   "}},
        {**base, "price": 1000},
    ):
        with pytest.raises(ValidationError):
            EoiDraftCreateRequest(**bad)


# ---- relevance notes ----------------------------------------------------------------------------------

REFERENCE = {
    "project_name": "Substation design Navoi", "client_name": "National Grid", "country": "Uzbekistan",
    "sector": "Energy", "service": "Engineering design", "role": "LEAD", "relevant_scope": "Detailed design of 220 kV substations",
    "start_date": "2019-01-01", "completion_date": "2021-06-30", "completion_state": "COMPLETED",
    "contract_value": "250000.00", "contract_currency": "USD",
}


def _request(key=1, language="en") -> NoteRequest:
    return NoteRequest(key=key, criterion_id="c-1", criterion_statement=SUBSTATION, reference=REFERENCE, language=language)


def test_relevance_note_validator_accepts_grounded_text_and_drops_the_rest() -> None:
    inputs = _request().input_text()
    good = "The Substation design Navoi assignment covered detailed design of 220 kV substations, completed in 2021."
    assert validate_note(good, inputs) == good
    assert validate_note("Work for National Grid in Uzbekistan worth 250,000 USD.", inputs)
    for bad in (
        "Completed in 2018 for National Grid.",                       # year not in inputs
        "A 300 kV design for National Grid.",                         # number not in inputs
        "Worth 250,000 EUR.",                                         # currency not in inputs
        "Designed for Asian Development Bank.",                        # name not in inputs
        "This shows the firm meets the criterion.",                   # compliance claim
        "Verified substation experience.",
        "One. Two. Three sentences are too many.",
        "",
        None,
        "x" * 600,
    ):
        assert validate_note(bad, inputs) is None, bad
    ru = "Задание Substation design Navoi для National Grid относится к проектированию подстанций."
    assert validate_note(ru, inputs) == ru
    assert validate_note("Задание подтверждает соответствие требованиям.", inputs) is None


def test_relevance_notes_are_bounded_and_failures_only_drop_notes() -> None:
    async def scenario():
        active = 0
        peak = 0

        async def call(request: NoteRequest) -> str:
            nonlocal active, peak
            active += 1
            peak = max(peak, active)
            await asyncio.sleep(0.05)
            active -= 1
            if request.key == 2:
                raise RuntimeError("provider error")
            if request.key == 3:
                return "Designed for Asian Development Bank."  # invalid: dropped
            if request.key == 4:
                await asyncio.sleep(1)  # beyond the per-call timeout
            return "The Substation design Navoi assignment covered detailed design of 220 kV substations."

        outcome = await generate_relevance_notes(
            [_request(key) for key in range(1, 9)], call, timeout_seconds=0.5, budget_seconds=5,
        )
        assert peak <= 4
        assert sorted(outcome.notes) == [1, 5, 6, 7, 8]
        assert (outcome.generated, outcome.dropped) == (5, 3)

        async def slow(request: NoteRequest) -> str:
            await asyncio.sleep(10)
            return "never"

        started = asyncio.get_running_loop().time()
        exhausted = await generate_relevance_notes([_request(key) for key in range(10)], slow, timeout_seconds=5, budget_seconds=0.3)
        assert asyncio.get_running_loop().time() - started < 2
        assert (exhausted.generated, exhausted.dropped) == (0, 10)
        assert (await generate_relevance_notes([], slow)).dropped == 0

    asyncio.run(scenario())


# ---- rendering ----------------------------------------------------------------------------------------

def _manifest(language: str) -> dict:
    return {
        "schema_version": "d2-05.eoi-manifest.v1", "language": language,
        "notice": {"assignment_title": {"value": "LOT-4 Detailed design SHINE project", "source": "SOURCE_TENDER"},
                   "reference_no": {"value": "OP00468882", "source": "SOURCE_TENDER"}},
        "letter": {"addressee_organization": "Ministry of Energy, Mongolia", "addressee_name": "Munkhbadral Purevsuren",
                   "signatory_name": "Aziza Karimova", "signatory_title": "Director", "contact_email": "office@example.uz",
                   "contact_phone": "+998 71 000 00 00", "contact_address": None, "source": "USER_INPUT"},
        "lead": {"firm_id": "f1", "display_name": "Codex Energy", "profile": [
            {"field": "name", "value": "Codex Energy", "source": {}}, {"field": "country", "value": "Uzbekistan", "source": {}},
            {"field": "sectors", "value": ["Energy"], "source": {}}, {"field": "legal_name", "value": None, "source": {}}]},
        "partners": [{"firm_id": "f2", "display_name": "Grid Partner LLP", "role": "JV_MEMBER",
                      "profile": [{"field": "name", "value": "Grid Partner LLP", "source": {}}]}],
        "experience": [
            {"no": 1, "reference_id": "r1", "firm_id": "f1", "firm_name": "Codex Energy", "eoi_role": "LEAD", "recorded_role": "LEAD",
             **{key: REFERENCE[key] for key in ("project_name", "client_name", "country", "start_date", "completion_date",
                                                "completion_state", "contract_value", "contract_currency", "relevant_scope")},
             "value_basis": "CONTRACT_TOTAL", "relevance_note": {"text": "Detailed design of 220 kV substations.", "criterion_id": "c1"}},
            {"no": 2, "reference_id": "r2", "firm_id": "f2", "firm_name": "Grid Partner LLP", "eoi_role": "JV_MEMBER",
             "recorded_role": "SUBCONSULTANT", "project_name": "Подстанция Алматы", "client_name": None, "country": "Kazakhstan",
             "start_date": None, "completion_date": None, "completion_state": "ONGOING", "contract_value": None,
             "contract_currency": None, "value_basis": "UNKNOWN", "relevant_scope": None, "relevance_note": None},
        ],
        "criteria": [
            {"requirement_id": "c1", "statement": SUBSTATION, "original_quote": SUBSTATION,
             "locator": {"page_number": None, "paragraph_number": 3}, "addressed_by": [1]},
            {"requirement_id": "c2", "statement": "Valid licenses", "original_quote": "The firm must hold valid licenses.",
             "locator": {"page_number": 2, "paragraph_number": None}, "addressed_by": []},
        ],
        "notes": [{"requirement_id": "n1", "note_kind": "SUBMISSION_INSTRUCTION", "statement": "Deliver by e-mail by 16 October 2026",
                   "original_quote": "Expressions of interest must be delivered by e-mail."}],
        "supporting_documents": [
            {"kind": "REFERENCE_PROOF", "row": 1, "name": "Substation design Navoi", "status": "ON_RECORD"},
            {"kind": "REFERENCE_PROOF", "row": 2, "name": "Подстанция Алматы", "status": "TO_ATTACH"},
            {"kind": "REGISTRATION", "name": None, "status": "TO_ATTACH"},
        ],
        "summary": {}, "price_information_included": False,
    }


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    parts.extend(paragraph.text for section in document.sections for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def _pdf_text(content: bytes) -> tuple[str, set[str]]:
    with fitz.open(stream=content, filetype="pdf") as document:
        text = "\n".join(page.get_text() for page in document)
        fonts = {font[3] for page in document for font in page.get_fonts()}
    return text, fonts


def test_documents_have_every_section_no_price_and_render_cyrillic() -> None:
    expectations = {
        "en": ["Expression of Interest", "Dear Munkhbadral Purevsuren,", "in association with Grid Partner LLP as joint-venture member",
               "Firm profile", "Relevant experience", "250,000 USD (contract total)", "Not recorded",
               "Shortlisting criteria cross-reference", "Addressed by references #1", "No selected reference addresses this criterion",
               "Notice, paragraph 3", "Notice, page 2", "Submission notes", "Supporting documents checklist", "On record",
               "To attach", "Draft note — verify", "Signature:", "No price information included."],
        "ru": ["Выражение заинтересованности", "совместно с Grid Partner LLP в качестве участника совместного предприятия",
               "Профиль фирмы", "Релевантный опыт", "250,000 USD (сумма контракта)", "Сопоставление с критериями отбора",
               "Относятся референции №1", "Объявление, абзац 3", "Подстанция Алматы", "Ценовая информация не включена."],
    }
    for language, needles in expectations.items():
        manifest = _manifest(language)
        docx_text = _docx_text(docx_bytes(manifest))
        pdf_text, fonts = _pdf_text(pdf_bytes(manifest))
        flat_pdf = " ".join(pdf_text.split())
        for needle in needles:
            assert needle in docx_text, (language, "docx", needle)
            assert " ".join(needle.split()) in flat_pdf, (language, "pdf", needle)
        assert any("Roboto" in font for font in fonts), fonts  # Cyrillic-capable font embedded
        for text in (docx_text, pdf_text):
            # The disclosure is the only place that may mention price, to say there is none.
            body = " ".join(text.split()).replace("No price information included.", "").replace("Ценовая информация не включена.", "")
            assert not PRICE_WORDS.search(body), PRICE_WORDS.search(body)
            for forbidden in ("compliant", "meets", "verified"):
                assert forbidden not in body.casefold().replace("verify:", "").replace("verify", ""), forbidden
        assert "None" not in docx_text  # a missing field is omitted or "Not recorded"
    assert docx_bytes(_manifest("en")) == docx_bytes(_manifest("en"))  # deterministic
    assert "Legal name" not in "\n".join(document_text(_manifest("en")))


# ---- migration ------------------------------------------------------------------------------------

def test_d2_05_migration_is_reversible_drift_free_and_immutable() -> None:
    async def scenario() -> None:
        database = support.database_name("d2_05_migration")
        await support.create_database(database)
        try:
            await support.raw_baseline(database)
            await asyncio.to_thread(support.alembic, database, "upgrade", D2_01_HEAD)
            await asyncio.to_thread(support.alembic, database, "upgrade", HEAD)
            connection = await support.database_connection(database)
            try:
                assert await connection.fetchval("SELECT version_num FROM alembic_version") == HEAD
                for table in ("eoi_drafts", "eoi_draft_artifacts"):
                    assert await connection.fetchval("SELECT to_regclass($1)", table) == table
                    assert await connection.fetchval(
                        "SELECT count(*) FROM pg_trigger WHERE tgname=$1", f"trg_{table}_immutable") == 1
            finally:
                await connection.close()
            await asyncio.to_thread(support.alembic, database, "downgrade", D2_01_HEAD)
            connection = await support.database_connection(database)
            try:
                assert await connection.fetchval("SELECT to_regclass('eoi_drafts')") is None
                assert await connection.fetchval("SELECT count(*) FROM pg_proc WHERE proname='reject_eoi_immutable_mutation'") == 0
            finally:
                await connection.close()
            await asyncio.to_thread(support.alembic, database, "upgrade", HEAD)
            check = await asyncio.to_thread(support.alembic, database, "check", success=False)
            assert check.returncode == 0, check.stderr or check.stdout
        finally:
            await support.drop_database(database)

    asyncio.run(scenario())


# ---- Postgres scenario ----------------------------------------------------------------------------

LETTER = {"addressee_organization": "Ministry of Energy", "addressee_name": None, "signatory_name": "Aziza Karimova",
          "signatory_title": "Director", "contact_email": "office@example.uz", "contact_phone": None, "contact_address": None}


def test_eoi_suggestions_drafts_isolation_immutability_and_staleness(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "PRIVATE_DOCUMENT_STORAGE_ROOT", str(tmp_path / "private"))

    async def no_provider(request):  # suggestions and validation never reach the provider
        raise AssertionError("the provider must not be called")

    monkeypatch.setattr(eoi_notes, "gemini_note_call", no_provider)

    async def scenario() -> None:
        database = support.database_name("d2_05_eoi")
        await support.create_database(database)
        try:
            await support.raw_baseline(database)
            await asyncio.to_thread(support.alembic, database, "upgrade", W1_HEAD)
            ids = await _seed_w1(database)
            await asyncio.to_thread(support.alembic, database, "upgrade", HEAD)
            connection = await support.database_connection(database)
            try:
                org_a = await connection.fetchval("SELECT id FROM organizations WHERE legacy_company_profile_id=$1", ids["profile_a"])
                org_b = await connection.fetchval("SELECT id FROM organizations WHERE legacy_company_profile_id=$1", ids["profile_b"])
                owner_a = await connection.fetchval("SELECT id FROM memberships WHERE organization_id=$1 AND user_id=$2", org_a, ids["user_a"])
                owner_b = await connection.fetchval("SELECT id FROM memberships WHERE organization_id=$1 AND user_id=$2", org_b, ids["user_b"])
            finally:
                await connection.close()
            engine = create_async_engine(support.target_url(database), pool_size=8)
            sessions = async_sessionmaker(engine, expire_on_commit=False)
            try:
                await _eoi_flow(sessions, database, ids, org_a, org_b, owner_a, owner_b, monkeypatch)
            finally:
                await engine.dispose()
        finally:
            await support.drop_database(database)

    asyncio.run(scenario())


async def _eoi_flow(sessions, database, ids, org_a, org_b, owner_a, owner_b, monkeypatch) -> None:
    issued = (
        f"{SUBSTATION}\n"
        "Key Experts will not be evaluated during the shortlisting stage.\n"
        "Expressions of interest must be delivered in written form via e-mail no later than 16 October 2026.\n"
        "The firm must hold valid licenses for high-complexity facility construction."
    )
    lines = issued.split("\n")
    async with sessions() as db:
        source = await get_or_create_source_pursuit(
            db, organization_id=org_a, actor_user_id=ids["user_a"], actor_membership_id=owner_a, tender_id=ids["tender"],
            stage=TenderEngagementStatus.SAVED, legacy_origin=TenderEngagementOrigin.OTHER_EXPLICIT_USER_ACTION,
        )
        pursuit_id = source.pursuit.id
        document = TenderDocument(
            tender_id=ids["tender"], file_url="d205.pdf", file_type="pdf", source_document_url="https://example.invalid/d205.pdf",
            source_document_type="RFP", sha256=hashlib.sha256(issued.encode()).hexdigest(), parsed_text=issued,
        )
        db.add(document)
        await db.commit()
        self_firm = await upsert_self_firm(db, organization_id=org_a, actor_user_id=ids["user_a"], payload=SelfFirmUpsertRequest(
            display_name="Codex Energy", country="Uzbekistan", sectors=["Energy"], services=["Engineering design"]))

        def reference(name, **values):
            base = dict(project_name=name, client_name="National Grid", country="Uzbekistan", sector="Energy",
                        service="Engineering design", role="LEAD", value_basis="UNKNOWN", start_date="2019-01-01",
                        completion_date="2021-06-30", completion_state="COMPLETED",
                        relevant_scope="Detailed design of 220 kV substations", evidence_state="UNVERIFIED")
            base.update(values)
            return ProjectReferenceCreateRequest(**base)

        navoi = await create_project_reference(db, organization_id=org_a, firm_id=self_firm.firm_id, actor_user_id=ids["user_a"],
            operator=False, payload=reference("Substation design Navoi", contract_value=Decimal("250000"), contract_currency="USD",
                                              value_basis="CONTRACT_TOTAL", evidence_provenance={"document_reference": "act.pdf"}))
        bukhara = await create_project_reference(db, organization_id=org_a, firm_id=self_firm.firm_id, actor_user_id=ids["user_a"],
            operator=False, payload=reference("Substation design Bukhara", completion_date="2023-03-01"))
        water = await create_project_reference(db, organization_id=org_a, firm_id=self_firm.firm_id, actor_user_id=ids["user_a"],
            operator=False, payload=reference("Water network design", sector="Water", service="Water supply design",
                                              relevant_scope="Design of water distribution networks"))
        partner = await create_firm(db, organization_id=org_a, actor_user_id=ids["user_a"], operator=False, payload=FirmCreateRequest(
            scope="ORGANIZATION_PRIVATE", canonical_name="Grid Partner", display_name="Grid Partner LLP", country="Kazakhstan",
            sectors=["Energy"], source_type="MANUAL", evidence_state="UNVERIFIED"))
        almaty = await create_project_reference(db, organization_id=org_a, firm_id=partner.firm_id, actor_user_id=ids["user_a"],
            operator=False, payload=reference("Almaty substation design", country="Kazakhstan", role="JV_MEMBER"))
        other = await create_firm(db, organization_id=org_a, actor_user_id=ids["user_a"], operator=False, payload=FirmCreateRequest(
            scope="ORGANIZATION_PRIVATE", canonical_name="Other", display_name="Other Firm", source_type="MANUAL",
            evidence_state="UNVERIFIED"))

        candidate = await build_analysis_pack_candidate(db, organization_id=org_a, pursuit_id=pursuit_id)
        started = await create_analysis_run(db, organization_id=org_a, pursuit_id=pursuit_id, membership_id=owner_a,
            request=PursuitAnalysisStartRequest(candidate_sha256=candidate.candidate_sha256, analysis_language="en",
                                                source_document_ids=[document.id]))

    async def extracted(sealed, language):
        item = sealed[0]

        def verified(fact):
            start = issued.index(fact.original_quote)
            return VerifiedFact(item.pack_item_id, fact, start, start + len(fact.original_quote), None, lines.index(fact.original_quote) + 1)

        return [
            verified(_fact(original_quote=lines[0], normalized_text=lines[0])),
            verified(_fact(original_quote=lines[1], normalized_text=lines[1], category="PERSONNEL",
                           requirement_type="EVALUATION_CRITERIA", stage_scope="SHORTLISTING", distinction="INFORMATIONAL", predicate=None)),
            verified(_fact(original_quote=lines[2], normalized_text="Deliver the expression of interest by e-mail by 16 October 2026",
                           category="SUBMISSION", requirement_type="SUBMISSION_INSTRUCTION", stage_scope="SUBMISSION", predicate=None)),
            verified(_fact(original_quote=lines[3], normalized_text=lines[3], category="LEGAL", requirement_type="CERTIFICATION",
                           predicate=None)),
        ]

    monkeypatch.setattr(analysis_service.pursuit_analyzer, "analyze_pack_items", extracted)
    async with sessions() as db:
        queued = await db.get(AnalysisRun, started.analysis_run_id)
        with pytest.raises(EoiConflictError, match="COMPLETED"):
            await eoi_suggestions(db, organization_id=org_a, pursuit_id=pursuit_id, analysis_run_id=queued.id)
        await process_analysis_run(db, started.analysis_run_id, worker_id="d2-05")
    run_id = started.analysis_run_id

    # ---- suggestions: passive, deterministic -------------------------------------------------------
    async with sessions() as db:
        counts_before = await _row_counts(db)
        suggestions = await eoi_suggestions(db, organization_id=org_a, pursuit_id=pursuit_id, analysis_run_id=run_id)
        assert await _row_counts(db) == counts_before
    assert suggestions.run_current is True
    assert suggestions.defaults.assignment_title == "W2 Source Tender"
    assert suggestions.defaults.reference_no == "WB-W2-1"
    assert suggestions.defaults.firm_name == "Codex Energy" and suggestions.defaults.firm_country == "Uzbekistan"
    assert suggestions.defaults.addressee_name == "Procurement Contact"
    assert len(suggestions.criteria) == 2 and len(suggestions.notes) == 2  # notes excluded from criteria
    experience = next(item for item in suggestions.criteria if item.original_quote == lines[0])
    assert experience.locator.paragraph_number == 1
    assert set(experience.matched_reference_ids) == {navoi.reference_id, bukhara.reference_id}
    own = {item.project_name: item for item in suggestions.own_references}
    assert [item.project_name for item in suggestions.own_references] == [
        "Substation design Bukhara", "Substation design Navoi", "Water network design",
    ]  # matched count desc, then completion recency desc
    assert [item.rank for item in suggestions.own_references] == [1, 2, 3]
    assert own["Substation design Navoi"].suggested and not own["Water network design"].suggested
    assert own["Substation design Navoi"].evidence_basis == "FILE_BACKED"
    firms = {item.display_name: item for item in suggestions.partner_firms}
    assert set(firms) == {"Grid Partner LLP", "Other Firm"}  # private partners only, never the own firm
    assert firms["Grid Partner LLP"].covers_requirement_ids == [experience.requirement_id]
    assert {note.note_kind for note in suggestions.notes} == {"INFORMATIONAL", "SUBMISSION_INSTRUCTION"}

    # ---- validation and isolation -------------------------------------------------------------------
    def request(**values):
        base = dict(analysis_run_id=run_id, language="en", own_reference_ids=[navoi.reference_id, bukhara.reference_id],
                    partners=[{"firm_id": partner.firm_id, "role": "JV_MEMBER", "reference_ids": [almaty.reference_id]}],
                    letter=LETTER, include_relevance_notes=False)
        base.update(values)
        return EoiDraftCreateRequest(**base)

    async with sessions() as db:
        for bad, message in (
            (request(own_reference_ids=[almaty.reference_id]), "own firm"),
            (request(partners=[{"firm_id": self_firm.firm_id, "role": "JV_MEMBER", "reference_ids": []}]), "own firm cannot be a partner"),
            (request(partners=[{"firm_id": other.firm_id, "role": "SUBCONSULTANT", "reference_ids": [almaty.reference_id]}]), "that firm"),
            (request(partners=[{"firm_id": uuid4(), "role": "SUBCONSULTANT", "reference_ids": []}]), "partner firm"),
        ):
            with pytest.raises(EoiConflictError, match=message):
                await create_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, membership_id=owner_a, request=bad)
        for call in (
            eoi_suggestions(db, organization_id=org_b, pursuit_id=pursuit_id, analysis_run_id=run_id),
            create_eoi_draft(db, organization_id=org_b, pursuit_id=pursuit_id, membership_id=owner_b, request=request()),
            eoi_suggestions(db, organization_id=org_a, pursuit_id=pursuit_id, analysis_run_id=uuid4()),
        ):
            with pytest.raises(EoiNotFoundError):
                await call
        assert await db.scalar(select(func.count(EoiDraft.id))) == 0

    # ---- a draft with relevance notes (mocked provider) ----------------------------------------------
    async def provider(note_request: NoteRequest) -> str:
        if note_request.reference["project_name"] == "Substation design Navoi":
            return "The Substation design Navoi assignment covered detailed design of 220 kV substations."
        return "Completed in 2015 for Asian Development Bank."  # ungrounded: dropped

    async with sessions() as db:
        draft = await create_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, membership_id=owner_a,
                                       request=request(include_relevance_notes=True), note_call=provider)
    assert draft.version == 1 and draft.current and draft.stale_reasons == []
    assert sorted(item.format for item in draft.artifacts) == ["DOCX", "PDF"]
    assert draft.summary.model_dump() == {
        "criteria_total": 2, "criteria_with_references": 1, "criteria_without_references": 1,
        "own_reference_count": 2, "partner_count": 1, "relevance_notes_generated": 1, "relevance_notes_dropped": 2,
    }
    async with sessions() as db:
        row = await db.get(EoiDraft, draft.draft_id)
        manifest = row.manifest
        assert manifest["price_information_included"] is False
        assert [item["source"]["id"] for item in manifest["experience"]] == [
            str(navoi.reference_id), str(bukhara.reference_id), str(almaty.reference_id)]
        assert all(item["source"]["type"] == "NOTICE_QUOTE" for item in manifest["criteria"] + manifest["notes"])
        assert manifest["experience"][0]["relevance_note"]["source"]["type"] == "GENERATED_DRAFT_NOTE"
        docx = next(item for item in draft.artifacts if item.format == "DOCX")
        artifact, _, path = await resolve_eoi_artifact(db, organization_id=org_a, pursuit_id=pursuit_id,
                                                       artifact_id=docx.artifact_id, membership_id=owner_a)
        text = _docx_text(path.read_bytes())
        assert "Draft note — verify: The Substation design Navoi assignment" in text
        assert "in association with Grid Partner LLP as joint-venture member" in text
        assert "Asian Development Bank" not in text
        assert "Addressed by references #1, #2, #3" in text
        assert hashlib.sha256(path.read_bytes()).hexdigest() == docx.sha256
        with pytest.raises(EoiNotFoundError):
            await resolve_eoi_artifact(db, organization_id=org_b, pursuit_id=pursuit_id, artifact_id=docx.artifact_id,
                                       membership_id=owner_b)
        assert await get_eoi_draft(db, organization_id=org_b, pursuit_id=pursuit_id, draft_id=draft.draft_id) is None

        ru = await create_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, membership_id=owner_a,
                                    request=request(language="ru"))
        assert ru.version == 2
        pdf = next(item for item in ru.artifacts if item.format == "PDF")
        _, _, pdf_path = await resolve_eoi_artifact(db, organization_id=org_a, pursuit_id=pursuit_id,
                                                    artifact_id=pdf.artifact_id, membership_id=owner_a)
        pdf_text, _ = _pdf_text(pdf_path.read_bytes())
        assert "Выражение заинтересованности" in pdf_text and "Релевантный опыт" in " ".join(pdf_text.split())

    # ---- immutability -------------------------------------------------------------------------------
    connection = await support.database_connection(database)
    try:
        for statement in ("UPDATE eoi_drafts SET language='ru' WHERE id=$1", "DELETE FROM eoi_drafts WHERE id=$1",
                          "UPDATE eoi_draft_artifacts SET byte_size=1 WHERE draft_id=$1",
                          "DELETE FROM eoi_draft_artifacts WHERE draft_id=$1"):
            with pytest.raises(Exception, match="immutable"):
                await connection.execute(statement, draft.draft_id)
    finally:
        await connection.close()

    # ---- staleness (projection only; stale drafts stay downloadable) -------------------------------
    async with sessions() as db:
        await update_project_reference(db, organization_id=org_a, firm_id=self_firm.firm_id, reference_id=navoi.reference_id,
                                       actor_user_id=ids["user_a"], payload=ProjectReferenceUpdateRequest(client_name="NEGU"), operator=False)
        assert (await get_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, draft_id=draft.draft_id)).stale_reasons == [
            "REFERENCE_SUPERSEDED"]
        await archive_project_reference(db, organization_id=org_a, firm_id=partner.firm_id, reference_id=almaty.reference_id,
                                        actor_user_id=ids["user_a"], operator=False)
        stale = await get_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, draft_id=draft.draft_id)
        assert stale.stale_reasons == ["REFERENCE_SUPERSEDED", "REFERENCE_ARCHIVED"] and not stale.current
        document_row = await db.scalar(select(TenderDocument).where(TenderDocument.file_url == "d205.pdf"))
        document_row.parsed_text = issued + "\nAddendum."
        await db.commit()
        stale = await get_eoi_draft(db, organization_id=org_a, pursuit_id=pursuit_id, draft_id=draft.draft_id)
        assert "ANALYSIS_INPUTS_CHANGED" in stale.stale_reasons
        document_row.parsed_text = issued
        await db.commit()
        candidate = await build_analysis_pack_candidate(db, organization_id=org_a, pursuit_id=pursuit_id)
        newer = await create_analysis_run(db, organization_id=org_a, pursuit_id=pursuit_id, membership_id=owner_a,
            request=PursuitAnalysisStartRequest(candidate_sha256=candidate.candidate_sha256, analysis_language="en",
                                                source_document_ids=[document_row.id]))
        await process_analysis_run(db, newer.analysis_run_id, worker_id="d2-05")
        drafts = await list_eoi_drafts(db, organization_id=org_a, pursuit_id=pursuit_id)
        assert [item.version for item in drafts] == [2, 1]
        assert "NEWER_ANALYSIS_RUN" in drafts[1].stale_reasons
        old = await eoi_suggestions(db, organization_id=org_a, pursuit_id=pursuit_id, analysis_run_id=run_id)
        assert old.run_current is False
        # The superseded Navoi reference is suggested under its current version.
        assert any(item.project_name == "Substation design Navoi" and item.client_name == "NEGU" for item in old.own_references)
        artifact, _, path = await resolve_eoi_artifact(db, organization_id=org_a, pursuit_id=pursuit_id,
                                                       artifact_id=docx.artifact_id, membership_id=owner_a)
        assert path.is_file()
        assert await db.scalar(select(func.count(EoiDraftArtifact.id))) == 4


async def _row_counts(db) -> dict[str, int]:
    from app.models.candidate_retrieval import Firm, ProjectReference
    from app.models.pursuit_analysis import AnalysisReviewAssertion
    counts = {}
    for model in (EoiDraft, EoiDraftArtifact, Firm, ProjectReference, AnalysisRun, AnalysisReviewAssertion):
        counts[model.__tablename__] = await db.scalar(select(func.count()).select_from(model))
    return counts
